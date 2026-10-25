"""Build the Assignment 3 report (.docx) following the official template structure.

Every number comes from results/*.json produced by the other scripts:
  results/availability.json   <- scripts/availability.py
  results/experiments.json    <- scripts/experiments.py   (Part D; marked PENDING if missing)
  results/dr_drill.json       <- scripts/dr_drill.py      (Part E evidence; optional)
Team data comes from team.yaml. Text still containing TODO/PENDING is printed in red.

Run: python scripts/build_report.py   ->  report/Assignment3_Report.docx
"""
import json
import subprocess
from datetime import date
from pathlib import Path

import yaml
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
RES, FIG = ROOT / "results", ROOT / "diagrams"
RED, GREY, NAVY = RGBColor(0xC0, 0x1F, 0x1F), RGBColor(0x55, 0x55, 0x55), RGBColor(0x1F, 0x3A, 0x5F)


def load(name):
    p = RES / name
    return json.loads(p.read_text()) if p.exists() else None


AV = load("availability.json")
EX = load("experiments.json")
DR = load("dr_drill.json")
TEAM = yaml.safe_load((ROOT / "team.yaml").read_text())
if AV is None:
    raise SystemExit("results/availability.json missing - run: python scripts/availability.py")


def git_commit():
    try:
        return subprocess.run(["git", "-C", str(ROOT), "rev-parse", "--short", "HEAD"],
                              capture_output=True, text=True, check=True).stdout.strip()
    except Exception:
        return "TODO-commit"


COMMIT = git_commit()
doc = Document()
st = doc.styles["Normal"]
st.font.name, st.font.size = "Calibri", Pt(10.5)
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2)
    s.top_margin = s.bottom_margin = Cm(1.8)


# ---------- helpers ----------
def run(p, text, bold=False, italic=False, size=None, mono=False, color=None):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if mono:
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    if color is not None:
        r.font.color.rgb = color
    elif "TODO" in text or "PENDING" in text:
        r.font.color.rgb = RED
    return r


def para(text="", bold=False, italic=False, size=None, align=None, space_after=4):
    p = doc.add_paragraph()
    if text:
        run(p, text, bold=bold, italic=italic, size=size)
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def rich(parts, space_after=4):
    """parts: list of (text, bold)"""
    p = doc.add_paragraph()
    for text, b in parts:
        run(p, text, bold=b)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def shade(cell, hex_fill):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def table(header, rows, widths=None, bold_first_col=True, size=9):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        run(c.paragraphs[0], h, bold=True, size=size)
        shade(c, "DCE6F1")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            run(cells[i].paragraphs[0], str(v), bold=(i == 0 and bold_first_col), size=size)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(path, caption, width=16.5):
    if path.exists():
        doc.add_picture(str(path), width=Cm(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        para(f"PENDING: {path.relative_to(ROOT)} not generated yet")
    para(caption, italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)


def code(text, size=7.5):
    for line in text.rstrip("\n").split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run(p, line if line else " ", mono=True, size=size, color=RGBColor(0x22, 0x22, 0x22))
    doc.add_paragraph()


def pc(x, d=4):
    return f"{100 * x:.{d}f}%"


H = doc.add_heading
B = lambda s: s  # noqa: E731  (readability)

# ---------- derived values ----------
t = AV["tiers"]
e2e = AV["end_to_end"]
share = AV["unavailability_share_percent"]
opt = AV["upgrade_options"]
gw2 = opt["2x API Gateway"]
lb2 = opt["2x Load Balancer"]
both = opt["2x API Gateway + 2x Load Balancer"]
mc = AV["monte_carlo_baseline"]
journ = AV["journeys"]
sweep = {s["change"].split(":")[0]: s for s in AV["plus_one_replica_sweep"]}
exps = {e["experiment"]: e for e in (EX or {}).get("experiments", [])}
members = TEAM["members"]


def fmt_s(v):
    return "n/a" if v is None else f"{v:.2f} s"


# =================== TITLE + SCENARIO (template header) ===================
p = para(align=WD_ALIGN_PARAGRAPH.CENTER)
run(p, "Assignment 3 — Cloud Computing and Microservices Resilience", bold=True, size=16, color=NAVY)
p = para(align=WD_ALIGN_PARAGRAPH.CENTER)
run(p, "AITU Digital Student Services Platform", bold=True, size=13)
para("Astana IT University | Fault Tolerance | Academic Year 2026–2027", italic=True,
     align=WD_ALIGN_PARAGRAPH.CENTER)
para("Context note: All operational data in this assignment are synthetic educational data. They are "
     "designed to represent a realistic Kazakhstan/Astana university context and do not describe actual "
     "AITU infrastructure, incidents, or internal systems.", italic=True, size=9)

H("Scenario", 1)
para("Design a resilient cloud-native architecture for a synthetic Astana IT University digital "
     "student-services platform. The platform supports timetable lookup, course information, registration "
     "requests, and student authentication.")
para("Architecture: Users → Load Balancer → API Gateway → Auth / Course / Registration → Database.")
H("Given requirements", 2)
table(["Parameter", "Requirement"], [
    ("Students/users", "10,000"), ("Peak requests", "500 req/s"), ("Availability target", "99.9%"),
    ("Maximum RTO", "15 min"), ("Maximum RPO", "5 min"), ("Maximum API latency", "500 ms"),
    ("Auth replicas", "2"), ("Course replicas", "2"), ("Registration replicas", "2"),
    ("Database nodes", "2")], widths=[7, 7])
H("Synthetic component availability", 2)
table(["Component", "Availability"], [
    ("Load Balancer", "99.95%"), ("API Gateway", "99.90%"), ("Auth instance", "99.50%"),
    ("Course instance", "99.50%"), ("Registration instance", "99.50%"), ("Database node", "99.70%")],
    widths=[7, 7])

# =================== PART A ===================
H("Part A — Architecture Analysis (5 points)", 1)
table(["Component", "Failure scenario", "Effect on student service", "Resilience mechanism"], [
    ("Load Balancer",
     "LB/ingress node lost, expired TLS certificate, or a wrong routing rule pushed to production.",
     "Nothing is reachable: login, timetable, course pages and enrolment all fail at the edge. In the "
     f"model it is an unreplicated block ({share['Load Balancer']:.1f}% of total unavailability).",
     "Managed regional LB spread over ≥2 zones (or 2 ingress replicas behind a floating IP), automated "
     "certificate renewal with expiry alerting, routing changes only through reviewed GitOps commits."),
    ("API Gateway",
     "Gateway process crash, worker/connection exhaustion when registration opens, or a slow upstream "
     "holding connections open.",
     "Every API answers 502/504 even though back-end pods are healthy. Largest single contributor: "
     f"{share['API Gateway']:.1f}% of modelled downtime.",
     "Deployed as 2 stateless NGINX replicas with a PodDisruptionBudget (minAvailable 1); 150 ms connect "
     "and 450 ms read timeouts keep the 500 ms budget; one retry on connection errors for GETs only."),
    ("Auth",
     "One replica OOM-killed or crashed; both replicas lost; identity store slow.",
     "One lost: logins continue on the survivor (it must be sized for the full peak). Both lost: new "
     "logins fail, but students already holding a valid short-lived token continue because tokens are "
     "verified at the gateway with cached keys.",
     "2 replicas on different nodes, readiness/liveness probes, signed JWTs with cached public keys, "
     "short token lifetime, timeouts below the gateway timeout."),
    ("Course",
     "Replica crash, slow catalogue/timetable query, both replicas lost.",
     "One lost: no visible effect. Both lost: timetable and course pages fail; enrolment still works "
     "because it does not call Course synchronously.",
     "2 replicas, short-TTL read cache with serve-stale-on-error (timetables change rarely), query "
     "timeout lower than the gateway timeout, no synchronous dependency from Registration."),
    ("Registration",
     "Process crash, hung process, DB timeouts, both replicas lost during the peak, duplicate submits "
     "after a client timeout.",
     "One lost: hidden by the Service, the survivor serves all traffic. Both lost: enrolment fails "
     "until pods return; Auth and Course are not affected. Without protection a retried click could "
     "double-book a seat.",
     "2-replica Deployment, startup/readiness/liveness probes, RollingUpdate maxUnavailable 0, PDB, "
     "400 ms DB timeout with immediate 503, Idempotency-Key + atomic seat update in one SQL statement."),
    ("Database",
     "Primary crash, disk full or corrupted, replication lag, accidental DELETE, both nodes lost.",
     "Primary lost: writes pause for the failover window and the services answer 503 quickly instead "
     "of hanging. A logical error (bad DELETE) is replicated to the standby instantly.",
     "Primary + hot standby with streaming replication and automatic failover (e.g. CloudNativePG or "
     "Patroni) with fencing; base backups + continuous WAL archive for point-in-time recovery (Part E)."),
], widths=[2.4, 4.3, 5.0, 5.3], size=8.5)
figure(FIG / "architecture.png",
       "Figure 1. Architecture of the synthetic platform as deployed in namespace aitu-a3 (solid) and as "
       "modelled (dashed). The orange gateway is the bottleneck that we replicated.")

# =================== PART B ===================
H("Part B — Availability Calculation (7 points)", 1)
para("For two independent replicas use Rparallel = 1 − (1 − R1)(1 − R2). Calculate each redundant "
     "subsystem and approximate end-to-end availability.")


def par_row(name, n, unit):
    a = t[name]["availability"]
    return (name, f"{n} {unit} (R = {AV['inputs'][name]['unit_availability']})",
            f"1 − (1 − {AV['inputs'][name]['unit_availability']})² = 1 − "
            f"{(1 - AV['inputs'][name]['unit_availability'])**2:.6f} = {a:.6f} ({pc(a)})")


table(["Subsystem", "Configuration", "Availability"], [
    par_row("Auth", 2, "replicas"), par_row("Course", 2, "replicas"),
    par_row("Registration", 2, "replicas"), par_row("Database", 2, "nodes"),
    ("End-to-end", "Series: LB × GW × Auth × Course × Registration × DB",
     f"0.9995 × 0.9990 × {t['Auth']['availability']:.6f}³ × {t['Database']['availability']:.6f} = "
     f"{e2e['availability']:.6f} ({pc(e2e['availability'])})")], widths=[2.8, 5.2, 9])
rich([("Does the design meet 99.9%? ", True),
      (f"No. The series product is {pc(e2e['availability'])}, {e2e['gap_pp']:.4f} percentage points "
       f"below the target. That is about {e2e['hours_per_year']} h of downtime per year "
       f"({e2e['minutes_per_month']} min per 30-day month) against an allowance of "
       f"{e2e['target_downtime']['hours_per_year']} h/yr ({e2e['target_downtime']['minutes_per_month']} "
       f"min/month). A Monte Carlo simulation with {mc['samples']:,} independent samples gives "
       f"{pc(mc['availability'])} (95% CI {pc(mc['ci95'][0])}–{pc(mc['ci95'][1])}), confirming the "
       "analytic result. Looking at individual student journeys does not rescue it: enrolment "
       f"(LB-GW-Auth-Registration-DB) is {pc(journ['Enroll in a course: LB-GW-Auth-Registration-DB'])}.",
       False)])
rich([("Which component is the bottleneck? ", True),
      (f"The single API Gateway. Its unavailability (0.0010) is {share['API Gateway']:.1f}% of the "
       f"total; the single Load Balancer adds {share['Load Balancer']:.1f}%, and all four replicated "
       f"tiers together only {100 - share['API Gateway'] - share['Load Balancer']:.1f}%. A 99.90% block "
       "in series uses up the entire 99.9% budget on its own, so no back-end improvement can compensate.",
       False)])
rich([("What redundancy change would have the largest effect? ", True),
      (f"A second API Gateway replica: the gateway tier rises to {pc(1 - 0.001**2)} and the system to "
       f"{pc(gw2['availability'])} (+{sweep['API Gateway']['gain_pp']:.4f} pp, {gw2['hours_per_year']} "
       f"h/yr), which meets the target. A second Load Balancer alone reaches only "
       f"{pc(lb2['availability'])}, and a third replica of Auth/Course/Registration/DB adds at most "
       f"{sweep['Auth']['gain_pp']:.4f} pp. Replicating both edge tiers gives {pc(both['availability'])} "
       f"({both['hours_per_year']} h/yr). Because the gateway is stateless, the change costs one small "
       "pod; we therefore deployed the gateway with 2 replicas in k8s/30-gateway.yaml.", False)])
figure(FIG / "rbd.png", "Figure 2a. Reliability block diagram with unit and tier availabilities.")
figure(FIG / "sensitivity.png",
       "Figure 2b. End-to-end availability for every single +1 replica change (scripts/availability.py).",
       width=14)

# =================== PART C ===================
H("Part C — Kubernetes Resilience Design (7 points)", 1)
para("Create Kubernetes YAML for the synthetic Registration Service. Do not use real AITU credentials, "
     "endpoints, or infrastructure.")
for item in ["2 replicas", "Readiness probe", "Liveness probe", "Resource requests and limits",
             "Rolling update", "Automatic restart", "Service discovery"]:
    doc.add_paragraph(item, style="List Bullet")
table(["Mechanism", "Implemented?", "Purpose"], [
    ("Replicas", "Yes — replicas: 2 (k8s/20-registration.yaml)",
     "Two active pods give the parallel availability used in Part B; preferred pod anti-affinity "
     "places them on different nodes when the cluster has more than one."),
    ("Liveness probe", "Yes — GET /livez, every 5 s, timeout 1 s, 3 failures",
     "Restarts a container whose process is stuck (deadlock). /livez never touches the database, so a "
     "DB outage cannot trigger restart loops. A startupProbe (30 × 1 s) protects cold starts."),
    ("Readiness probe", "Yes — GET /ready, every 2 s, timeout 1 s, 2 failures",
     "Checks that the pod can reach PostgreSQL within 400 ms and is not shutting down. A failing pod is "
     "removed from the Service endpoints so the gateway stops sending it students."),
    ("Resource limits", "Yes — requests 150m / 96Mi, limits 500m / 192Mi",
     "Requests reserve capacity at scheduling; limits contain a memory leak or CPU spike in one pod so "
     "it cannot starve its neighbours (OOM-kill affects only that container)."),
    ("Rolling update", "Yes — maxSurge 1, maxUnavailable 0, minReadySeconds 3, preStop sleep 3 s",
     "A new pod must be Ready (and stay Ready 3 s) before an old one is removed; the preStop delay plus "
     "the service's SIGTERM drain lets in-flight requests finish, so deploys cause no errors."),
    ("Automatic restart", "Yes — restartPolicy: Always + ReplicaSet + PDB minAvailable 1",
     "Kubelet restarts crashed containers in place; the ReplicaSet re-creates deleted pods; the PDB "
     "stops node drains from evicting both replicas at once."),
    ("Service discovery", "Yes — ClusterIP Service + CoreDNS",
     "Gateway and clients use registration-service.aitu-a3.svc.cluster.local; only Ready pods are "
     "behind the virtual IP. The service finds the DB the same way (registration-db.aitu-a3.svc...)."),
], widths=[3, 5, 9], size=8.5)
para(f"Part C artifact — k8s/20-registration.yaml (commit {COMMIT}); Deployment excerpt:", bold=True)
yml = (ROOT / "k8s" / "20-registration.yaml").read_text()
start = yml.index("kind: Deployment")
end = yml.index("---", start)
code("apiVersion: apps/v1\n" + yml[start:end])

# =================== PART D ===================
H("Part D — Failure Simulation (6 points)", 1)
para("Run/simulate: one Registration replica failure, all Registration replicas failing, and database "
     "unavailability.")


def exp_row(n, label):
    e = exps.get(n)
    if not e:
        return (str(n), label, "PENDING — run make experiments", "PENDING", "PENDING")
    impact = (f"{e['requests_failed']}/{e['requests_total']} requests failed "
              f"({e['success_percent']}% success), error window {e['error_window_s']} s; "
              f"p95 of successful requests {e['p95_ms_ok']} ms")
    if e["requests_failed"]:
        kinds = ", ".join(f"{v}× {'no response' if k == '0' else 'HTTP ' + k}"
                          for k, v in e["failed_by_status"].items())
        impact += f"; failures: {kinds}, returned in p95 {e['p95_ms_failed']} ms"
    rec = fmt_s(e["recovery_s"])
    if e.get("recovery_after_restore_s") is not None:
        rec += f" from injection ({e['recovery_after_restore_s']:.2f} s after DB restored)"
    extra = f"; restarts {e['restarts_before']}→{e['restarts_after']}"
    if e["pods_recreated"]:
        extra += f"; new pods {', '.join(p[-5:] for p in e['pods_recreated'])}"
    return (str(n), e["injection"], fmt_s(e["detection_s"]), rec + extra, impact)


table(["Experiment", "Injected failure", "Detection time", "Recovery time", "User impact"], [
    exp_row(1, "One Registration replica"), exp_row(2, "All Registration replicas"),
    exp_row(3, "Database unavailable")] + ([exp_row(4, "One replica hangs")] if 4 in exps else []),
    widths=[1.6, 3.6, 2.2, 4.2, 5.4], size=8)
if EX:
    para(f"Method. Measured on Kubernetes {EX['kubernetes_server']} ({EX['run_at']}). An in-cluster pod "
         f"(student-traffic) sent {EX['load']['rate_rps']:.0f} req/s (70% course list, 30% enrolment) "
         "through the api-gateway Service for "
         f"{EX['load']['duration_s']:.0f} s; the fault was injected after {EX['load']['inject_at_s']:.0f} s. "
         "Detection = time from injection until Kubernetes reported fewer than 2 Ready registration pods "
         "(polled every ~0.2 s). Recovery = time until 2 pods were Ready again. User impact is computed "
         "from every individual request logged by the generator (results/raw/). The pod clock was "
         "aligned to the host clock before the run.", size=9.5)
else:
    para("PENDING: run `make experiments` on your cluster, then rebuild this report — the table and "
         "Figure 3 are filled in automatically from results/experiments.json.")
figure(FIG / "experiments.png",
       "Figure 3. Per-request latency/status (top) and Ready registration pods (bottom); dashed black = "
       "fault injected, dashed green = database restored.", width=17)
log_path = RES / "experiment_log.txt"
if log_path.exists():
    para("Extract from results/experiment_log.txt (kubectl get pods during the faults):", bold=True)
    lines = log_path.read_text().splitlines()
    picks = []
    for i, l in enumerate(lines):
        if "DURING (t+1 s)" in l:
            picks += lines[i:i + 6]
    code("\n".join(picks[:24]), size=6.5)

# =================== PART E ===================
H("Part E — RPO/RTO and Disaster Recovery (5 points)", 1)
drill = ""
if DR:
    drill = (f" Lab drill (results/dr_drill.json): full logical restore of the registration DB "
             f"measured at {DR['restore_seconds']:.1f} s, service Ready again after "
             f"{DR['service_ready_seconds']:.1f} s, rows verified {DR['rows_after']}/{DR['rows_before']}.")
table(["Mechanism", "Frequency", "Expected RPO", "Recovery time", "Purpose"], [
    ("Database backup",
     "Daily base backup 01:30 (low traffic) + WAL segment archived at least every 60 s "
     "(archive_timeout = 60) to object storage outside the cluster; 14-day retention.",
     "≤ 60 s of committed data (age of the last archived WAL) — meets ≤ 5 min.",
     "≈ 6–10 min: fetch base backup, replay WAL to the chosen timestamp (PITR)." + drill,
     "Recover from what replication cannot: accidental DELETE/DROP, bad migration, ransomware, loss "
     "of both DB nodes."),
    ("Replica",
     "Continuous asynchronous streaming replication primary → hot standby; lag alert at 10 s.",
     "Seconds (normally < 1 s of lag); 0 if switched to synchronous commit for registration week.",
     "10–30 s automatic promotion of the standby; < 10 s for app pods (Part D).",
     "Survive a single DB node, disk or process failure without restoring anything."),
    ("Configuration backup",
     "Every change: manifests in Git (this repository); secrets in an external secret store; nightly "
     "etcd snapshot of the cluster.",
     "0 for manifests (Git is the source of truth); ≤ 24 h for ad-hoc cluster state.",
     "≈ 2–3 min: `kubectl apply -f k8s/` to a fresh namespace or cluster.",
     "Rebuild deployments, probes, services and PDBs exactly, with no manual drift."),
    ("Full recovery",
     "On disaster (cluster/zone lost); rehearsed once per semester, before registration week.",
     "≤ 60 s (bounded by the WAL archive).",
     "Target ≤ 14 min: new cluster/namespace 3 min + PITR restore 8 min + pods Ready and smoke test "
     "2 min + DNS/LB switch 1 min — inside the 15 min RTO.",
     "Bring the whole student-services platform back in another failure domain when HA is defeated."),
], widths=[2.5, 4.0, 3.2, 3.8, 3.5], size=8)
para("Explain why high availability and disaster recovery are different engineering objectives.", bold=True)
para("High availability is about not noticing failures. It assumes faults are local and independent — one "
     "pod crashes, one node reboots, the DB primary dies — and hides them while the system keeps running, "
     "using spare live capacity: two replicas, probes that route around bad pods, a standby that takes "
     "over. It is measured continuously as availability (99.9%) and works in seconds. Its weakness is that "
     "every copy is live and connected: a DELETE without WHERE, a broken migration or a lost data centre "
     "reaches the primary and the standby at the same moment, so HA faithfully keeps serving wrong data or "
     "nothing at all.")
para("Disaster recovery starts from the assumption that HA has already failed. Its tools are copies that "
     "are deliberately separated in time and place — backups and archived WAL in another location, and "
     "configuration in Git — plus a rehearsed runbook. It is not measured by uptime but by two limits "
     "that apply after the disaster: how much committed data may be lost (RPO ≤ 5 min) and how long "
     "restoration may take (RTO ≤ 15 min). A system can therefore have excellent HA and no DR (fast "
     "failover, but nothing to restore from after corruption) or good DR and poor HA (data is safe, but "
     "every crash is an outage); the platform needs both, and they are tested separately — HA in Part D, "
     "DR in the restore drill.")

# =================== SUBMISSION ===================
H("Submission", 1)
for item in ["Architecture diagram", "Availability calculations", "Kubernetes YAML",
             "Failure experiment results", "RPO/RTO plan", "Oral defense"]:
    doc.add_paragraph(item, style="List Bullet")
doc.add_page_break()

# =================== SUBMISSION RECORD ===================
p = para(align=WD_ALIGN_PARAGRAPH.CENTER)
run(p, "STUDENT TECHNICAL WORK REPORT / SUBMISSION RECORD", bold=True, size=13, color=NAVY)
para("Assignment 3 — Cloud Computing and Microservices Resilience", bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER)
para("Purpose: This form is a formal submission record. The student must complete every applicable field "
     "and attach evidence of their own work. Short generic explanations are not sufficient.",
     italic=True, size=9)

files = ("app/registration/service.py, app/registration/Dockerfile, app/loadgen/loadgen.py, "
         "k8s/00-namespace.yaml, k8s/10-postgres.yaml, k8s/20-registration.yaml, k8s/30-gateway.yaml, "
         "k8s/40-loadgen.yaml, scripts/availability.py, scripts/experiments.py, scripts/dr_drill.py, "
         "scripts/diagrams.py, scripts/build_report.py, results/*.json, results/experiment_log.txt, "
         "docs/dr-plan.md")
H("1. Student and Submission Information", 1)
table(["Required field", "Student entry"], [
    ("Student full name", ", ".join(m["name"] for m in members)),
    ("Student ID", ", ".join(str(m["id"]) for m in members)),
    ("Group / cohort", TEAM["group"]), ("Instructor", TEAM["instructor"]),
    ("Date of submission", TEAM["submission_date"]),
    ("Individual variant / task code assigned by instructor", TEAM["variant"]),
    ("Git repository / project link", f"{TEAM['repo_url']} (commit {COMMIT})"),
    ("Attached code / file names", files)], widths=[6, 11])

H("2. Work Period and Task Record", 1)
table(["Item", "Student entry"], [
    ("Work start date/time", TEAM["work_start"]), ("Work completion date/time", TEAM["work_end"]),
    ("Main task completed", "Parts A–E: failure analysis of all six tiers, series/parallel availability "
     "model with Monte Carlo check, Registration Service on Kubernetes with full probe/rollout/PDB "
     "configuration, three measured failure experiments (plus a liveness experiment), RPO/RTO plan."),
    ("Tools / technologies used", "Kubernetes " + (EX["kubernetes_server"] if EX else "(version from run)")
     + ", Docker, kubectl, PostgreSQL 16, NGINX 1.27, Python 3.12 (stdlib http.server, pg8000, numpy, "
     "matplotlib, python-docx), Git/GitHub"),
    ("Main technical output produced", "Runnable Registration service + gateway + DB manifests, automated "
     "failure-injection harness with raw per-request logs, availability model, figures, this report "
     "(generated from the result files).")], widths=[6, 11])

H("3. Report Structure to Complete", 1)
para("For each section below, write the requested evidence in your own words. Include calculations, "
     "diagrams, code references, screenshots, logs, or experimental evidence where applicable.", size=9)


def record(title, intro, summary, evidence, decision):
    H(title, 2)
    para(intro, size=9)
    table(["Required content", "Student response / evidence"], [
        ("Summary", summary), ("Evidence / calculation / artifact reference", evidence),
        ("Student-specific decision or observation", decision)], widths=[5, 12])


e1, e2, e3 = exps.get(1), exps.get(2), exps.get(3)
record("1. Architecture Analysis",
       "Analyze the failure scenarios for each component and explain the resilience mechanism selected for "
       "your design.",
       "Each of the six tiers was analysed for single-instance faults, loss of all replicas and slow "
       "dependencies. Edge tiers (LB, gateway) are protected by replication and strict timeouts, stateless "
       "services by 2 replicas with probes and safe rollouts, and the database by a hot standby plus "
       "backups for faults that replication would copy.",
       "Part A table; Figure 1 (diagrams/architecture.png); k8s/30-gateway.yaml (timeouts, retry policy).",
       "We separated 'slow' from 'down': every hop has a timeout shorter than the hop before it "
       "(DB 400 ms < gateway 450 ms < 500 ms SLA), so a hanging database produces a fast 503 instead of "
       "a pile-up of waiting gateway connections.")
record("2. Availability Calculation",
       "Show the parallel-redundancy calculations and the end-to-end series calculation. Explain whether "
       "the target is met and identify the bottleneck.",
       f"Parallel tiers: Auth/Course/Registration {pc(t['Auth']['availability'])}, Database "
       f"{pc(t['Database']['availability'])}. Series with LB and gateway: {pc(e2e['availability'])} — "
       f"target missed ({e2e['hours_per_year']} h/yr vs {e2e['target_downtime']['hours_per_year']} h/yr). "
       f"Bottleneck: API Gateway ({share['API Gateway']:.1f}% of unavailability).",
       "scripts/availability.py → results/availability.json; Figures 2a and 2b.",
       f"We verified the closed-form result with a seeded Monte Carlo run ({pc(mc['availability'])}) and "
       f"swept every +1 replica option: only the gateway change reaches the target "
       f"({pc(gw2['availability'])}); a third back-end replica adds ≤ {sweep['Auth']['gain_pp']:.4f} pp.")
record("3. Kubernetes Resilience Design",
       "Include the Registration Service YAML or a clear excerpt/reference. Explain replicas, probes, "
       "resources, rolling update, restart, and service discovery.",
       "Registration runs as a 2-replica Deployment with startup, readiness (DB-aware) and liveness "
       "(process-only) probes, requests/limits, RollingUpdate maxSurge 1/maxUnavailable 0 with "
       "minReadySeconds 3, preStop drain, restartPolicy Always, a PodDisruptionBudget and a ClusterIP "
       "Service; runs as non-root with a read-only filesystem.",
       f"k8s/20-registration.yaml, app/registration/service.py (commit {COMMIT}); Part C table and excerpt.",
       "Readiness result is cached for 0.5 s in the service so that probes plus traffic do not open a DB "
       "connection storm, and SIGTERM first flips readiness to false and keeps serving for 5 s — combined "
       "with the 3 s preStop this removes errors during rolling updates.")
if EX:
    sum4 = (f"Exp 1 (crash of one replica): detected {fmt_s(e1['detection_s'])}, recovered "
            f"{fmt_s(e1['recovery_s'])}, {e1['requests_failed']}/{e1['requests_total']} failed. "
            if e1 else "")
    sum4 += (f"Exp 2 (both pods deleted): detected {fmt_s(e2['detection_s'])}, recovered "
             f"{fmt_s(e2['recovery_s'])}, {e2['requests_failed']}/{e2['requests_total']} failed. "
             if e2 else "")
    sum4 += (f"Exp 3 (DB down): detected {fmt_s(e3['detection_s'])}, recovered "
             f"{fmt_s(e3['recovery_after_restore_s'])} after DB restore, restarts "
             f"{e3['restarts_before']}→{e3['restarts_after']}." if e3 else "")
else:
    sum4 = "PENDING — filled automatically after `make experiments`."
record("4. Failure Simulation",
       "Document all three required simulations with detection time, recovery time, and user impact. "
       "Explain how each result was obtained.",
       sum4,
       "scripts/experiments.py; results/experiments.json; results/raw/exp*_requests.jsonl (every request); "
       "results/raw/exp*_timeline.jsonl (pod states); results/experiment_log.txt; Figure 3.",
       "Experiments 1 and 2 deliberately test two different self-healing paths: a process crash is fixed "
       "by the kubelet restarting the container in the same pod, while deleting both pods is fixed by the "
       "ReplicaSet creating new ones. The crash is triggered from inside the pod through a localhost-only "
       "endpoint, because `kill -9 1` from kubectl exec is ignored for a container's PID 1.")
record("5. RPO/RTO and Disaster Recovery",
       "Complete the RPO/RTO plan and explain why high availability and disaster recovery are different "
       "engineering objectives.",
       "Backups (daily base + WAL archived ≤ 60 s) give RPO ≤ 1 min and a PITR restore of ≈ 6–10 min; "
       "streaming replication covers node loss in 10–30 s; Git-managed manifests rebuild the cluster in "
       "minutes; the full-recovery runbook totals ≤ 14 min, within the 15 min RTO.",
       "Part E table, docs/dr-plan.md (runbook)" + (", results/dr_drill.json (measured restore)." if DR
                                                    else "; restore drill: scripts/dr_drill.py."),
       "Backups are stored outside the cluster and restored into an empty database during the drill — a "
       "backup that has never been restored is not counted as a recovery mechanism.")
record("6. Final Architecture Decision",
       "State the most important resilience change you made and justify it using availability, failure, "
       "or recovery evidence.",
       "Run the API Gateway as two stateless replicas (with a PDB) in front of the 2-replica Registration "
       "service, combined with DB-aware readiness and process-only liveness.",
       f"Availability {pc(e2e['availability'])} → {pc(gw2['availability'])} "
       f"({e2e['hours_per_year']} → {gw2['hours_per_year']} h/yr) in results/availability.json"
       + (f"; Exp 3 showed {e3['restarts_before']}→{e3['restarts_after']} restarts during the DB outage."
          if e3 else "."),
       "In a series system the weakest unreplicated block decides the result; adding more back-end "
       "replicas is almost worthless until the gateway is duplicated.")

H("4. Individual Work and Anti-Plagiarism Record", 1)
para("Because the assignment uses a common educational scenario and dataset, the submitted report must "
     "contain verifiable evidence of the individual student's work. Copying another student's wording, "
     "calculations, architecture, code, logs, screenshots, or conclusions is not acceptable.", size=9)
for item in ["Use the instructor-assigned individual variant/task code.",
             "Provide the exact YAML filename or repository commit used for the reported configuration.",
             "Include evidence from your own simulation or test run.",
             "Explain one Kubernetes setting you changed or tuned and why.",
             "Be prepared during oral defense to explain what happens if both replicas fail."]:
    doc.add_paragraph(item, style="List Bullet")
impl = (f"Measured: " + "; ".join(
    f"Exp {n} detect {fmt_s(exps[n]['detection_s'])} / recover {fmt_s(exps[n]['recovery_s'])}"
    for n in sorted(exps)) if exps else "PENDING experiment numbers")
table(["Individual evidence", "Student entry"], [
    ("My unique technical decision",
     "Timeout chain DB 400 ms < gateway 450 ms < SLA 500 ms with fail-fast 503, and readiness turned "
     "off on SIGTERM + preStop 3 s so rolling updates and DB outages never leave requests hanging. "
     "Tuned setting: readinessProbe periodSeconds 2 / failureThreshold 2, so a pod that loses the DB "
     "leaves the Service within ~4 s, while liveness (15 s) never reacts to the DB at all."),
    ("My own implementation / calculation evidence",
     f"scripts/availability.py ({pc(e2e['availability'])} baseline, {pc(gw2['availability'])} with 2× "
     f"gateway, Monte Carlo {pc(mc['availability'])}); {impl}."),
    ("Exact file, code, commit, notebook, or log reference",
     f"{TEAM['repo_url']} commit {COMMIT}: k8s/20-registration.yaml, app/registration/service.py, "
     "scripts/experiments.py, results/experiments.json, results/experiment_log.txt"),
    ("One problem I personally encountered and how I solved it", TEAM["individual"]["problem_encountered"]),
    ("One limitation of my own solution", TEAM["individual"]["limitation"]),
], widths=[5, 12])
table(["Team member", "Student ID", "GitHub", "Own contribution (see git log)"],
      [(m["name"], str(m["id"]), str(m["github"]), m["owns"]) for m in members], widths=[3.5, 2, 2.5, 9])

H("5. Evidence Checklist", 1)
checks = [
    ("Completed required tables and calculations", bool(EX)),
    ("Architecture / RBD / FTA / security diagram where applicable", (FIG / "rbd.png").exists()),
    ("Source code or configuration files where applicable", True),
    ("Execution logs / experiment results where applicable", bool(EX)),
    ("Screenshots or other reproducible evidence where applicable", (FIG / "experiments.png").exists()),
    ("Links to code, datasets, notebooks, or artifacts", "TODO" not in TEAM["repo_url"]),
    ("Before/after comparison where applicable", True),
    ("Final technical conclusion", True)]
for text, ok in checks:
    para(("☑ " if ok else "☐ ") + text)

H("6. Student Declaration", 1)
para("I confirm that this report describes my own work. I have not copied another student's report, code, "
     "calculations, diagrams, logs, screenshots, or conclusions. Any external source or reused material has "
     "been identified. I understand that I may be asked to reproduce or explain any part of this "
     "submission during the oral defense.", size=9.5)
para("AI assistance: parts of the code and text were drafted with an AI assistant (Claude) and then run, "
     "checked and edited by the team; all measurements come from the team's own cluster run.",
     italic=True, size=9)
table(["Role", "Name", "Signature", "Date"],
      [("Student / Executor", f"{m['name']} ({m['id']})", "", TEAM["submission_date"]) for m in members]
      + [("Instructor / Accepted by", TEAM["instructor"], "________________", "____________")],
      widths=[4, 6, 3.5, 3.5])

H("7. Instructor Acceptance Record", 1)
table(["Check", "Status / comment"], [
    ("Report structure completed", "☐ Yes  ☐ No  Comment: ______________________________"),
    ("Evidence is sufficient", "☐ Yes  ☐ No  Comment: ______________________________"),
    ("Individual contribution is identifiable", "☐ Yes  ☐ No  Comment: ______________________________"),
    ("Oral defense completed", "☐ Yes  ☐ No  Comment: ______________________________"),
    ("Final status", "☐ Accepted  ☐ Revision required")], widths=[6, 11])
para("Fault Tolerance | Astana IT University context", italic=True, size=8,
     align=WD_ALIGN_PARAGRAPH.CENTER)

out = ROOT / "report" / "Assignment3_Report.docx"
out.parent.mkdir(exist_ok=True)
doc.save(out)
todo = [m["name"] for m in members if "TODO" in m["name"]]
print(f"wrote {out.relative_to(ROOT)}  (experiments: {'yes' if EX else 'PENDING'}, "
      f"dr drill: {'yes' if DR else 'no'}, TODO team fields: {'yes' if todo else 'no'})")
